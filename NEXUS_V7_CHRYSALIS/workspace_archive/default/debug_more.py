import fitz
from docx import Document
import os

BASE_DIR = r"C:\Code\NEXUS\20_NEXUS - Copie\external_sources\DOCADI-MAS"

def debug_more():
    # DOCX
    path_docx = os.path.join(BASE_DIR, "Templates", "Template.docx")
    print(f"--- DOCX Search 'ADI' ---")
    doc = Document(path_docx)
    found = False
    for p in doc.paragraphs:
        if "ADI" in p.text:
            print(f"Found 'ADI' in Body: '{p.text}'")
            found = True
    for section in doc.sections:
        for p in section.header.paragraphs:
            if "ADI" in p.text:
                print(f"Found 'ADI' in Header: '{p.text}'")
                found = True
    if not found: 
        print("'ADI' not found in DOCX text.")

    # PDF
    path_pdf = os.path.join(BASE_DIR, "Source", "PR4_MAB.pdf")
    print(f"\n--- PDF Search 'ADI' ---")
    doc = fitz.open(path_pdf)
    page = doc[0]
    text = page.get_text()
    if "ADI" in text:
        print("Found 'ADI' in PDF text.")
    else:
        print("'ADI' not found in PDF text.")
    
    print("\n--- PDF Image Analysis ---")
    images = page.get_images(full=True)
    print(f"Found {len(images)} images on page 1.")
    for img in images:
        xref = img[0]
        rects = page.get_image_rects(xref)
        for r in rects:
            print(f"Image rect: {r}")

if __name__ == "__main__":
    debug_more()
