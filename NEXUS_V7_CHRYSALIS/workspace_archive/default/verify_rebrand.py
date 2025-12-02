import os
import fitz  # PyMuPDF
from docx import Document

def verify_pdf(file_path):
    print(f"Checking PDF: {file_path}")
    if not os.path.exists(file_path):
        print("  [FAIL] File not found.")
        return False
    
    try:
        doc = fitz.open(file_path)
        text = ""
        for page in doc:
            text += page.get_text()
        
        errors = 0
        if "AD Industries" in text:
            print("  [FAIL] Found 'AD Industries'")
            errors += 1
        if "ADI" in text: 
            # Note: "ADI" might be a substring of other words, so simple check might have false positives.
            # But for this verification, let's check generic presence or try to be smarter if needed.
            # Let's assume strict replacement was requested.
            # Actually, "ADI" is short, might match "TRADITION". Let's check "ADI " or similar if possible, 
            # but the requirements were vague. Let's stick to exact match for now but warn.
            print("  [WARNING] Found 'ADI' (could be substring)")
            
        if "Motherson Aerospace" not in text:
            print("  [FAIL] Missing 'Motherson Aerospace'")
            errors += 1
        # MAS might also be a substring
        if "MAS" not in text:
            print("  [FAIL] Missing 'MAS'")
            errors += 1
            
        if errors == 0:
            print("  [PASS] Text verification successful.")
            return True
        return False
    except Exception as e:
        print(f"  [ERROR] {e}")
        return False

def verify_docx(file_path):
    print(f"Checking DOCX: {file_path}")
    if not os.path.exists(file_path):
        print("  [FAIL] File not found.")
        return False

    try:
        doc = Document(file_path)
        text = ""
        for para in doc.paragraphs:
            text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + "\n"
        for section in doc.sections:
            for para in section.header.paragraphs:
                text += para.text + "\n"
            for para in section.footer.paragraphs:
                text += para.text + "\n"

        errors = 0
        if "AD Industries" in text:
            print("  [FAIL] Found 'AD Industries'")
            errors += 1
        if "Motherson Aerospace" not in text:
            print("  [FAIL] Missing 'Motherson Aerospace'")
            errors += 1
            
        if errors == 0:
            print("  [PASS] Text verification successful.")
            return True
        return False
    except Exception as e:
        print(f"  [ERROR] {e}")
        return False

if __name__ == "__main__":
    base_dir = os.getcwd()
    
    artifacts = [
        "Rebranded_Template.docx",
        "Rebranded_PR4-HY_Rev_02_Fiche_processus_INDUSTRIALISER.pdf",
        "Rebranded_PR4-MCO_Rev_03_Fiche Processus Industrialiser.pdf",
        "Rebranded_PR4_MAB.pdf"
    ]
    
    print("=== VERIFICATION REPORT ===")
    for artifact in artifacts:
        path = os.path.join(base_dir, artifact)
        if artifact.endswith(".docx"):
            verify_docx(path)
        elif artifact.endswith(".pdf"):
            verify_pdf(path)
    print("===========================")
