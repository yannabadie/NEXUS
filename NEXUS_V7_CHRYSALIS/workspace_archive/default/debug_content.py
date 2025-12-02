import fitz
from docx import Document
import os

BASE_DIR = r"C:\Code\NEXUS\20_NEXUS - Copie\external_sources\DOCADI-MAS"

def debug_docx():
    path = os.path.join(BASE_DIR, "Templates", "Template.docx")
    print(f"--- Debugging DOCX: {path} ---")
    doc = Document(path)
    found = False
    for p in doc.paragraphs:
        if "AD Industries" in p.text:
            print(f"Found in paragraph: '{p.text}'")
            found = True
            # Check runs
            for i, run in enumerate(p.runs):
                print(f"  Run {i}: '{run.text}'")
    
    # Check headers
    for section in doc.sections:
        for p in section.header.paragraphs:
             if "AD Industries" in p.text:
                print(f"Found in Header: '{p.text}'")
                found = True
    
    if not found:
        print("String 'AD Industries' NOT found in DOCX (text checks).")

def debug_pdf():
    path = os.path.join(BASE_DIR, "Source", "PR4_MAB.pdf") # Check one PDF
    print(f"\n--- Debugging PDF: {path} ---")
    doc = fitz.open(path)
    page = doc[0]
    text = page.get_text()
    if "AD Industries" in text:
        print("Found 'AD Industries' in PDF text.")
    else:
        print("String 'AD Industries' NOT found in PDF text. Dumping first 500 chars:")
        print(text[:500])

if __name__ == "__main__":
    debug_docx()
    debug_pdf()
